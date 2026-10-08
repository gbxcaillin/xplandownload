"""Find client folders that are the same person under two Xplan records, and merge them.

Xplan often has the same client twice ("Alexander Mc Donough (40975)" and "Alexander
Mcdonough (41479)"). Folders are grouped when their names match ignoring spaces, case and
punctuation (or the same words in another order), or when the clients share a surname and
date of birth. Each extra folder is then checked against the group's main folder using
Xplan:

* conflicts (keep separate): different dates of birth, different middle names, a person vs a
  trust/company, or the two records are each other's partner;
* strong evidence: same date of birth, email, phone or address;
* supporting evidence: the same Xplan documents filed under both records.

``find_duplicates`` writes a review workbook with a Decision per folder ("merge" or "keep",
pre-filled; change it where needed). ``apply_merges`` merges the folders marked "merge":
files move into the main folder; a file that is the very same Xplan document as one already
there is removed (to the SharePoint recycle bin); a different file with the same name is
renamed "<name> (from <id>)". documents_index.csv follows every move (old copy kept).
"""

from __future__ import annotations

import csv
import datetime as dt
import os
import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Iterable

from .documents import _long
from .folders import GROUPS, ID_SUFFIX, INDEX, _client_folder, _folders, _parts

Progress = Callable[[str], None]
REVIEW_PREFIX = "folder_merge_review_"
DECISION = "Decision (merge / keep)"
HEADS = ["Group", "Main folder", "Main in", "Main Xplan IDs", "Merge in folder", "Merge in in",
         "Merge in Xplan IDs", "Files (main / merge in)", "Same documents in both",
         "Evidence", "Conflicts", "Confidence", DECISION, "Notes"]


# --------------------------------------------------------------------------
# What we know about each Xplan record
# --------------------------------------------------------------------------

@dataclass
class Person:
    id: int
    type: str = ""
    first: str = ""
    middle: str = ""
    last: str = ""
    dob: str = ""
    emails: set[str] = field(default_factory=set)
    phones: set[str] = field(default_factory=set)
    addresses: set[str] = field(default_factory=set)
    partner: int | None = None


def _letters(text) -> str:
    return re.sub(r"[^a-z]", "", str(text or "").lower())


def _dob(value) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, (dt.date, dt.datetime)):
        return value.strftime("%Y-%m-%d")
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", str(value))
    if m:
        return "-".join(m.groups())
    m = re.match(r"(\d{1,2})/(\d{1,2})/(\d{4})", str(value))
    return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}" if m else ""


def _phone(value) -> str:
    digits = re.sub(r"\D", "", str(value or ""))
    return digits[-9:] if len(digits) >= 8 else ""


def _address(street, postcode) -> str:
    s = re.sub(r"[^a-z0-9]", "", str(street or "").lower())
    p = re.sub(r"\D", "", str(postcode or ""))
    return f"{s}|{p}" if s and p else ""


def people_from_xplan(data, wanted: Iterable[int]) -> dict[int, Person]:
    """Person facts for the wanted ids, from a loaded XplanData."""
    out: dict[int, Person] = {}
    for eid in wanted:
        e = data.entities.get(eid)
        p = Person(eid)
        if e:
            f = e.f
            p.type = e.type
            p.first = _letters(f.get("first_name") or f.get("preferred_name"))
            p.middle = _letters(f.get("middle_name"))
            p.last = _letters(f.get("last_name"))
            p.dob = _dob(f.get("dob"))
            if "@" in str(f.get("preferred_email") or ""):
                p.emails.add(str(f["preferred_email"]).strip().lower())
            if _phone(f.get("preferred_phone")):
                p.phones.add(_phone(f.get("preferred_phone")))
            a = _address(f.get("preferred_street"), f.get("preferred_postcode"))
            if a:
                p.addresses.add(a)
        for r in data.contacts.get(eid, []):
            value, kind = str(r.get("value") or "").strip(), str(r.get("type") or "").lower()
            if "fax" in kind or not value:
                continue
            if "@" in value:
                p.emails.add(value.lower())
            elif _phone(value):
                p.phones.add(_phone(value))
        for r in data.addresses.get(eid, []):
            a = _address(r.get("street"), r.get("postcode"))
            if a:
                p.addresses.add(a)
        p.partner = data.partner_of.get(eid)
        out[eid] = p
    return out


# --------------------------------------------------------------------------
# Comparing two folders
# --------------------------------------------------------------------------

def compare(a_ids: Iterable[int], b_ids: Iterable[int], people: dict[int, Person],
            shared_docs: int) -> tuple[list[str], list[str]]:
    """(evidence, conflicts) between the Xplan records behind two folders."""
    evidence: list[str] = []
    conflicts: list[str] = []

    def add(bucket, text):
        if text not in bucket:
            bucket.append(text)

    for ai in a_ids:
        for bi in b_ids:
            a, b = people.get(ai), people.get(bi)
            if not a or not b or ai == bi:
                continue
            if a.partner == bi or b.partner == ai:
                add(conflicts, "the two records are partners")
            if a.type and b.type and (a.type == "individual") != (b.type == "individual"):
                add(conflicts, f"one is a {a.type}, the other a {b.type}")
            if a.dob and b.dob:
                if a.dob == b.dob:
                    add(evidence, "same date of birth")
                else:
                    add(conflicts, "different dates of birth")
            if a.first and b.first and not (a.first.startswith(b.first[:3])
                                            or b.first.startswith(a.first[:3])):
                add(conflicts, "different first names")
            if a.middle and b.middle and a.middle[0] != b.middle[0]:
                add(conflicts, "different middle names")
            if a.emails & b.emails:
                add(evidence, "same email")
            if a.phones & b.phones:
                add(evidence, "same phone")
            if a.addresses & b.addresses:
                add(evidence, "same address")
    if shared_docs:
        add(evidence, f"{shared_docs} Xplan document(s) filed under both")
    return evidence, conflicts


def confidence(evidence: list[str], conflicts: list[str]) -> tuple[str, str]:
    """(confidence label, pre-filled decision)."""
    if conflicts:
        return "Different people?", "keep"
    strong = [e for e in evidence if e.startswith("same")]
    if strong:
        return "High", "merge"
    if evidence:
        return "Medium", "merge"
    return "Name only", "merge"


# --------------------------------------------------------------------------
# Finding the groups
# --------------------------------------------------------------------------

@dataclass
class Folder:
    name: str
    where: str            # "", "Active" or "Inactive"
    ids: set[int]
    items: set[tuple]     # Xplan document keys in it (from the index)
    files: int

    @property
    def label(self) -> str:
        return ID_SUFFIX.sub("", self.name).strip()


def name_keys(label: str) -> list[str]:
    words = re.findall(r"[a-z]+", label.lower())
    if not words:
        return []
    return [f"n:{''.join(words)}", f"s:{' '.join(sorted(words))}"]


def read_index(index_path: Path) -> tuple[dict[tuple[str, str], set[int]],
                                          dict[tuple[str, str], set[tuple]],
                                          dict[tuple[str, str], int]]:
    ids: dict[tuple[str, str], set[int]] = defaultdict(set)
    items: dict[tuple[str, str], set[tuple]] = defaultdict(set)
    files: dict[tuple[str, str], int] = defaultdict(int)
    if not index_path.is_file():
        return ids, items, files
    with open(_long(index_path), newline="", encoding="utf-8-sig") as fh:
        for row in csv.DictReader(fh):
            where = _client_folder(_parts(row.get("saved_as", "")))
            if not where:
                continue
            key = (where[0], where[1].lower())
            cid = (row.get("client_id") or "").strip()
            if cid.isdigit():
                ids[key].add(int(cid))
            items[key].add(_item_key(row))
            files[key] += 1
    return ids, items, files


def _item_key(row: dict) -> tuple:
    return (row.get("kind", ""), row.get("xplan_docid", ""), row.get("xplan_partid", ""))


def load_folders(dest: Path) -> list[Folder]:
    ids, items, files = read_index(dest / INDEX)
    out = []
    for where, path in _folders(dest / "Clients"):
        key = (where, path.name.lower())
        fid = set(ids.get(key, ()))
        if not fid:
            m = ID_SUFFIX.search(path.name)
            if m:
                fid = {int(m.group(1))}
        out.append(Folder(path.name, where, fid, items.get(key, set()), files.get(key, 0)))
    return out


def group_folders(folders: list[Folder], people: dict[int, Person] | None = None
                  ) -> list[list[Folder]]:
    """Folders that look like the same client, 2 or more per group."""
    parent = list(range(len(folders)))

    def find(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    seen: dict[str, int] = {}
    for i, f in enumerate(folders):
        keys = name_keys(f.label)
        for eid in f.ids:
            p = (people or {}).get(eid)
            if p and p.type == "individual" and p.dob and p.last:
                keys.append(f"d:{p.last}|{p.dob}")
        for k in keys:
            if k in seen:
                parent[find(i)] = find(seen[k])
            else:
                seen[k] = i
    groups: dict[int, list[Folder]] = defaultdict(list)
    for i, f in enumerate(folders):
        groups[find(i)].append(f)
    return [g for g in groups.values() if len(g) > 1]


def main_folder(group: list[Folder]) -> Folder:
    """Active first, then the most files, then the oldest Xplan record."""
    return min(group, key=lambda f: (f.where != "Active", -f.files, min(f.ids or {10 ** 12})))


@dataclass
class Proposal:
    group: int
    main: Folder
    other: Folder
    shared: int
    evidence: list[str]
    conflicts: list[str]
    level: str
    decision: str


def propose(groups: list[list[Folder]], people: dict[int, Person]) -> list[Proposal]:
    out = []
    for n, g in enumerate(sorted(groups, key=lambda g: main_folder(g).label.lower()), 1):
        main = main_folder(g)
        for f in sorted(g, key=lambda f: f.name.lower()):
            if f is main:
                continue
            shared = len(main.items & f.items)
            evidence, conflicts = compare(main.ids, f.ids, people, shared)
            level, decision = confidence(evidence, conflicts)
            out.append(Proposal(n, main, f, shared, evidence, conflicts, level, decision))
    return out


def write_review(proposals: list[Proposal], path: Path) -> None:
    import xlsxwriter

    wb = xlsxwriter.Workbook(str(path), {"strings_to_formulas": False, "strings_to_urls": False})
    bold = wb.add_format({"bold": True, "bg_color": "#DDEBF7", "border": 1, "text_wrap": True})
    fill = wb.add_format({"bg_color": "#FFF2CC", "border": 1})
    warn = wb.add_format({"font_color": "#9C0006"})
    ws = wb.add_worksheet("Merge review")
    for i, h in enumerate(HEADS):
        ws.write(0, i, h, bold)
    for r, p in enumerate(proposals, 1):
        ws.write_number(r, 0, p.group)
        ws.write(r, 1, p.main.name)
        ws.write(r, 2, p.main.where or "Clients")
        ws.write(r, 3, ", ".join(map(str, sorted(p.main.ids))))
        ws.write(r, 4, p.other.name)
        ws.write(r, 5, p.other.where or "Clients")
        ws.write(r, 6, ", ".join(map(str, sorted(p.other.ids))))
        ws.write(r, 7, f"{p.main.files} / {p.other.files}")
        ws.write_number(r, 8, p.shared)
        ws.write(r, 9, "; ".join(p.evidence) or "name only")
        ws.write(r, 10, "; ".join(p.conflicts), warn)
        ws.write(r, 11, p.level)
        ws.write_string(r, 12, p.decision, fill)
        ws.write_blank(r, 13, None)
    widths = [7, 40, 9, 14, 40, 9, 14, 12, 10, 44, 34, 16, 16, 30]
    for i, w in enumerate(widths):
        ws.set_column(i, i, w)
    ws.freeze_panes(1, 2)
    ws.autofilter(0, 0, max(len(proposals), 1), len(HEADS) - 1)
    ws.write(len(proposals) + 2, 0,
             "Each row merges the 'Merge in' folder into the 'Main' folder. Change the yellow "
             "Decision to 'keep' for any that are different people (rows with Conflicts are "
             "already 'keep'). Then run: python -m xplan_extract merge-folders --apply")
    wb.close()


def find_duplicates(dest: Path, data, progress: Progress = print) -> tuple[list[Proposal], Path]:
    """data: a loaded XplanData (or None to compare on names and documents only)."""
    folders = load_folders(dest)
    progress(f"Client folders: {len(folders):,}")
    all_ids = {i for f in folders for i in f.ids}
    people = people_from_xplan(data, all_ids) if data is not None else {}
    groups = group_folders(folders, people)
    proposals = propose(groups, people)
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    path = dest / f"{REVIEW_PREFIX}{stamp}.xlsx"
    write_review(proposals, path)
    return proposals, path


# --------------------------------------------------------------------------
# Merging
# --------------------------------------------------------------------------

def latest_review(dest: Path) -> Path | None:
    found = sorted(dest.glob(f"{REVIEW_PREFIX}*.xlsx"))
    return found[-1] if found else None


def read_decisions(path: Path) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Merge review"] if "Merge review" in wb.sheetnames else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    head = [str(h or "").strip() for h in next(rows, [])]
    need = ["Main folder", "Main in", "Merge in folder", "Merge in in", DECISION]
    if any(h not in head for h in need):
        raise ValueError(f"{path.name} isn't a merge review sheet")
    out = []
    for r in rows:
        row = {h: ("" if v is None else str(v).strip()) for h, v in zip(head, r)}
        if row.get("Main folder") and row.get("Merge in folder"):
            out.append(row)
    wb.close()
    return out


@dataclass
class MergeResult:
    merged: int = 0
    kept: int = 0
    skipped: int = 0        # left for later (not in the chosen confidence levels)
    files_moved: int = 0
    duplicates_removed: int = 0
    renamed: int = 0
    index_rows_updated: int = 0
    problems: list[str] = field(default_factory=list)
    log_file: Path | None = None


def _where(text: str) -> str:
    return text if text in GROUPS else ""


def _unique(folder: Path, name: str, hint: str) -> str:
    stem, ext = os.path.splitext(name)
    candidate = f"{stem} (from {hint}){ext}"
    n = 2
    while (folder / candidate).exists():
        candidate = f"{stem} (from {hint}-{n}){ext}"
        n += 1
    return candidate


def apply_merges(dest: Path, review: Path, progress: Progress = print,
                 only: Iterable[str] | None = None) -> MergeResult:
    """Merge the rows marked "merge"; with ``only``, just those Confidence levels."""
    levels = {x.strip().lower() for x in only} if only else None
    clients = dest / "Clients"
    index_path = dest / INDEX
    head: list[str] = []
    rows: list[list[str]] = []
    if index_path.is_file():
        with open(_long(index_path), newline="", encoding="utf-8-sig") as fh:
            reader = csv.reader(fh)
            head = next(reader, [])
            rows = list(reader)
    col = {h: i for i, h in enumerate(head)}
    si = col.get("saved_as")
    sep = "\\" if rows and si is not None and "\\" in rows[0][si] else "/"
    by_path: dict[str, list[int]] = defaultdict(list)   # lower saved_as -> row numbers
    for n, r in enumerate(rows):
        if si is not None and len(r) > si:
            by_path[r[si].replace("/", "\\").lower()].append(n)

    def rel(path: Path) -> str:
        return str(path.relative_to(dest)).replace("/", "\\").lower()

    def set_path(old: Path, new: Path) -> int:
        nums = by_path.pop(rel(old), [])
        new_rel = sep.join(new.relative_to(dest).parts)
        for n in nums:
            rows[n][si] = new_rel
            if "status" in col and len(rows[n]) > col["status"]:
                where = _client_folder(list(new.relative_to(dest).parts))
                rows[n][col["status"]] = where[0] if where else ""
        by_path[rel(new)].extend(nums)
        return len(nums)

    def item_of(path: Path) -> tuple | None:
        nums = by_path.get(rel(path))
        if not nums:
            return None
        r = rows[nums[0]]
        return _item_key({h: r[i] if i < len(r) else "" for h, i in col.items()})

    result = MergeResult()
    log: list[list[str]] = []
    for d in read_decisions(review):
        decision = d.get(DECISION, "").lower()
        if levels is not None and d.get("Confidence", "").strip().lower() not in levels:
            result.skipped += 1
            continue
        if not decision.startswith("merge"):
            result.kept += 1
            continue
        main_dir = clients / _where(d["Main in"]) / d["Main folder"]
        src_dir = clients / _where(d["Merge in in"]) / d["Merge in folder"]
        if not src_dir.is_dir():
            result.problems.append(f"{d['Merge in folder']}: folder not found (already merged?)")
            continue
        if not main_dir.is_dir():
            result.problems.append(f"{d['Main folder']}: main folder not found")
            continue
        hint = (d.get("Merge in Xplan IDs") or "").split(",")[0].strip() or "merged"
        main_items: dict[tuple, Path] = {}
        for p in main_dir.rglob("*"):
            if p.is_file() and item_of(p) is not None:
                main_items.setdefault(item_of(p), p)
        moved = dupes = renamed = 0
        for src in sorted(p for p in src_dir.rglob("*") if p.is_file()):
            target_dir = main_dir / src.parent.relative_to(src_dir)
            item = item_of(src)
            try:
                twin = main_items.get(item) if item is not None else None
                if twin is not None and twin.exists():
                    # the very same Xplan document is already in the main folder
                    os.remove(_long(src))
                    result.index_rows_updated += set_path(src, twin)
                    dupes += 1
                    continue
                target_dir.mkdir(parents=True, exist_ok=True)
                name = src.name
                if (target_dir / name).exists():
                    name = _unique(target_dir, name, hint)
                    renamed += 1
                os.replace(_long(src), _long(target_dir / name))
                result.index_rows_updated += set_path(src, target_dir / name)
                if item is not None:
                    main_items.setdefault(item, target_dir / name)
                moved += 1
            except OSError as exc:
                result.problems.append(f"{d['Merge in folder']}: {src.name}: "
                                       f"{type(exc).__name__} {exc.strerror or exc}")
        for sub in sorted((p for p in src_dir.rglob("*") if p.is_dir()), reverse=True):
            try:
                sub.rmdir()
            except OSError:
                pass
        try:
            src_dir.rmdir()
            result.merged += 1
            status = "merged"
            if result.merged % 50 == 0:
                progress(f"  merged {result.merged:,} folders")
        except OSError:
            status = "partly merged (files left behind: see problems)"
            result.problems.append(f"{d['Merge in folder']}: not empty after the merge")
        result.files_moved += moved
        result.duplicates_removed += dupes
        result.renamed += renamed
        log.append([d["Main folder"], d["Main in"], d["Merge in folder"], d["Merge in in"],
                    d.get("Main Xplan IDs", ""), d.get("Merge in Xplan IDs", ""), status,
                    str(moved), str(dupes), str(renamed)])

    if result.index_rows_updated:
        stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
        os.replace(_long(index_path),
                   _long(index_path.with_name(f"{index_path.stem}.before-merge-{stamp}.csv")))
        tmp = index_path.with_suffix(".tmp")
        with open(_long(tmp), "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh)
            w.writerow(head)
            w.writerows(rows)
        os.replace(_long(tmp), _long(index_path))
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    result.log_file = dest / f"folder_merge_done_{stamp}.csv"
    with open(_long(result.log_file), "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(["main_folder", "main_in", "merged_folder", "merged_from", "main_xplan_ids",
                    "merged_xplan_ids", "result", "files_moved", "duplicates_removed",
                    "renamed"])
        w.writerows(log)
    return result
