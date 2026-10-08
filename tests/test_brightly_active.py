import openpyxl

from xplan_extract.brightly.active import norm_name, read_active_list

HEAD = [None, None, "Transaction No", None, "Transaction  Date", "Policy Number", "Client No",
        "Client / Owner", "Product Name", "Type", "Transaction Type", None, "Adviser", None,
        "Amount", "GST", "Inc GST", "CRM Reference"]


def test_read_active_list(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([None, None, None, None, None, None, "Report title"])
    ws.append(HEAD)
    ws.append([None, None, "Clients: Citizen, Jane"])
    ws.append([None, None, "123", None, "2026-01-01", "PF-001 23", "9001", "Citizen, Jane",
               "Wrap", "FEE", "Ongoing", None, "Alex Adviser", None, 1, 0, 1, "101"])
    ws.append([None, None, "Total"])
    ws.append([None, None, "Clients: Single, Bob"])
    ws.append([None, None, "124", None, "2026-01-01", "777", "9002", "Single, Bob",
               "Wrap", "FEE", "Ongoing", None, "Alex Adviser", None, 1, 0, 1, None])
    path = tmp_path / "list.xlsx"
    wb.save(path)
    listed = read_active_list(path)
    assert [c.name for c in listed] == ["Citizen, Jane", "Single, Bob"]
    assert listed[0].crm_refs == {101} and listed[0].policies == {"PF00123"}
    assert listed[1].crm_refs == set() and listed[1].policies == {"777"}


def test_norm_name():
    assert norm_name("Citizen, Jane") == norm_name("Mrs Jane Citizen") == "citizen jane"


def test_name_variants_and_loose():
    from xplan_extract.brightly.active import loose_name, name_variants
    v = name_variants("Citizen, Jane & Sam")
    assert "Jane Citizen" in v and "Sam Citizen" in v
    v = name_variants("Jane & Sam Citizen")
    assert "Jane Citizen" in v and "Sam Citizen" in v
    v = name_variants("Example Pty Ltd ATF Example Family Trust")
    assert "Example Pty Ltd" in v and "Example Family Trust" in v
    assert loose_name("Example Super Fund Pty Ltd") == loose_name("EXAMPLE SMSF") == "example"


def test_excel_numbers(tmp_path):
    from xplan_extract.brightly.active import _cell
    assert _cell(12345.0) == "12345" and _cell(1.5) == "1.5" and _cell(None) == ""


def test_initial_keys():
    from xplan_extract.brightly.active import initial_keys, name_variants
    assert initial_keys("J Citizen") == initial_keys("Jane Citizen") == ["citizen|j"]
    assert initial_keys("Citizen, J A") == ["citizen|j"]
    assert initial_keys("CITIZEN J") == ["citizen|j"]
    assert initial_keys("Mrs J. Citizen") == ["citizen|j"]
    assert "Sam Citizen" in name_variants("Jane Citizen & Sam Citizen")


def test_parse_answer():
    from xplan_extract.brightly.active import parse_answer
    assert parse_answer("101, 102") == ["101", "102"]
    assert parse_answer("not in Xplan") == [] and parse_answer("N/A") == []
    assert parse_answer("") is None and parse_answer("ask Jo") is None


def test_confirm_workbook_keeps_answers(tmp_path):
    from types import SimpleNamespace

    from xplan_extract.brightly.active import (ListedClient, read_answers, read_confirmed,
                                               write_confirm_workbook)

    person = SimpleNamespace(f={"last_name": "Citizen"}, name="Jane Citizen", id=101)
    builder = SimpleNamespace(
        households={"hh-1": SimpleNamespace(people=[person],
                                            record_entity=SimpleNamespace(name="Citizen"))},
        d=SimpleNamespace(entities={}))
    a = ListedClient("Citizen, J", fees=10)
    b = ListedClient("Nobody, Ann", fees=5)
    c = ListedClient("Maybe, Max", fees=1)
    path = tmp_path / "confirm.xlsx"
    assert write_confirm_workbook([a, b, c], builder, {}, path) == 3

    wb = openpyxl.load_workbook(path)
    ws = wb["To confirm"]
    head = [h.value for h in ws[1]]
    for row in ws.iter_rows(min_row=2):
        name = row[0].value
        if name == "Citizen, J":
            row[head.index("Confirmed Xplan ID")].value = "101"
            row[head.index("Notes")].value = "goes by Jo"
        elif name == "Nobody, Ann":
            row[head.index("Confirmed Xplan ID")].value = "not in Xplan"
    wb.save(path)
    assert read_confirmed(path) == {"Citizen, J": ["101"], "Nobody, Ann": []}

    # refresh: Citizen is now matched (from the answer) but stays on the list, answered
    a.households, a.method = {"hh-1"}, "Confirmed by Scott"
    b.method = "Confirmed: not in Xplan"
    answers = read_answers(path)
    out = tmp_path / "confirm2.xlsx"
    assert write_confirm_workbook([a, b, c], builder, {}, out, answers) == 1
    ws = openpyxl.load_workbook(out)["To confirm"]
    head = [h.value for h in ws[1]]
    rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True) if r[1] is not None}
    assert set(rows) == {"Citizen, J", "Nobody, Ann", "Maybe, Max"}
    status = head.index("Status")
    assert rows["Citizen, J"][head.index("Notes")] == "goes by Jo"
    assert rows["Citizen, J"][status].startswith("Answered: matched to 1 family group")
    assert rows["Nobody, Ann"][status].startswith("Answered: not in Xplan")
    assert rows["Maybe, Max"][status] == "Waiting for an answer"
    assert read_confirmed(out) == {"Citizen, J": ["101"], "Nobody, Ann": []}
