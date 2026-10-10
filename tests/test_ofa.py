"""Ongoing fee arrangements: the consent rules, the database steps and the form.
Fictional data only."""

import datetime as dt

import pytest
import sqlalchemy as sa

from xplan_extract.brightly import database as db
from xplan_extract.brightly import loader, ofa, ofa_store
from xplan_extract.brightly.ofa_form import Practice, build, form_for

from test_brightly_loader import HOUSEHOLD, write_export

D = dt.date


def test_anniversary_and_window():
    assert ofa.anniversary(D(2020, 2, 29), 2026) == D(2026, 2, 28)
    w = ofa.window(D(2026, 6, 1))
    assert (w.opens, w.closes) == (D(2026, 4, 2), D(2026, 10, 29))


def test_business_days_skip_weekends_and_victorian_holidays():
    # Wed 29 Oct 2025 + 10 business days, with Melbourne Cup (Tue 4 Nov) in between
    assert ofa.add_business_days(D(2025, 10, 29), 10) == D(2025, 11, 13)
    hols = ofa.victorian_public_holidays(2026)
    assert D(2026, 4, 3) in hols and D(2026, 4, 6) in hols      # Good Friday, Easter Monday
    assert D(2026, 1, 26) in hols and D(2026, 6, 8) in hols     # Australia Day, King's Birthday
    hols27 = ofa.victorian_public_holidays(2027)
    assert D(2027, 12, 27) in hols27 and D(2027, 12, 28) in hols27  # Christmas on a Saturday


def test_status_through_the_year():
    ref = D(2019, 6, 1)
    st = lambda today, consents=(), **kw: ofa.status(ref, list(consents), today, **kw)
    assert st(D(2026, 1, 10), [D(2025, 6, 1)]).state == "covered"
    assert st(D(2026, 3, 10), [D(2025, 6, 1)]).state == "upcoming"
    assert st(D(2026, 4, 10), [D(2025, 6, 1)]).state == "open"
    assert st(D(2026, 9, 30), [D(2025, 6, 1)]).state == "due_soon"
    s = st(D(2026, 10, 20), [D(2025, 6, 1)])
    assert s.state == "urgent" and s.days_left == 9 and s.due_on == D(2026, 10, 29)
    s = st(D(2026, 11, 2), [D(2025, 6, 1)])
    assert s.state == "lapsed" and s.due_on == D(2026, 11, 13)   # 10 business days, Cup Day
    assert st(D(2026, 11, 2), [D(2025, 6, 1)], providers_notified=True).state == "ended"
    assert st(D(2026, 11, 2), [D(2025, 6, 1), D(2026, 6, 1)]).state == "covered"
    # a later consent shows the arrangement carried on, even if an earlier one isn't recorded
    assert st(D(2026, 11, 2), [D(2026, 6, 1)]).state == "covered"
    # Xplan history unknown: past windows are not called lapsed
    assert st(D(2026, 11, 2), consent_history_known=False).state in ("covered", "upcoming")
    assert ofa.status(None, [], D(2026, 1, 1)).state == "check"


def test_consent_anniversary_and_deadlines():
    ref = D(2019, 6, 1)
    assert ofa.consent_anniversary(ref, D(2026, 4, 2)) == D(2026, 6, 1)     # first day open
    assert ofa.consent_anniversary(ref, D(2026, 10, 29)) == D(2026, 6, 1)   # last day
    assert ofa.consent_anniversary(ref, D(2026, 10, 30)) is None
    assert ofa.consent_anniversary(ref, D(2026, 1, 15)) is None
    due = ofa.withdrawal_deadlines(D(2026, 12, 22))
    assert due["acknowledge_by"] == D(2027, 1, 8)    # skips Christmas, Boxing Day, New Year
    assert ofa.keep_until(D(2026, 7, 1)) == D(2031, 7, 1)


@pytest.fixture
def engine(tmp_path):
    e = sa.create_engine(f"sqlite:///{tmp_path / 'b.db'}")
    db.create_all(e)
    loader.load_export(e, write_export(tmp_path / "x"), progress=lambda m: None)
    return e


def test_seed_record_and_worklist(engine, tmp_path):
    out = ofa_store.seed(engine, "test", today=D(2026, 10, 10), dry_run=True)
    assert out["created"] == 1
    with engine.connect() as c:
        assert c.execute(sa.select(sa.func.count()).select_from(db.ofa_arrangement)).scalar() == 0
    out = ofa_store.seed(engine, "test", today=D(2026, 10, 10))
    assert out["created"] == 1 and ofa_store.seed(engine, "t")["skipped_existing"] == 1
    aid = f"OFA-{HOUSEHOLD['id']}-1"
    items = ofa_store.worklist(engine, D(2026, 10, 10))
    assert items[0]["status"].problems == ["no deduction account recorded",
                                           "last consent not confirmed (from Xplan)"]
    # HOUSEHOLD's OFA anniversary is 1 March: 2026 window closed 29 July, history unknown
    with pytest.raises(ofa_store.OfaError, match="verbal"):
        ofa_store.record_consent(engine, aid, D(2026, 3, 1), "verbal", "t")
    with pytest.raises(ofa_store.OfaError, match="outside"):
        ofa_store.record_consent(engine, aid, D(2026, 9, 1), "esign", "t")
    ofa_store.confirm_history(engine, aid, "pat", D(2026, 2, 20))
    ofa_store.add_account(engine, aid, "Sample Wrap", "pat", account_number="W2",
                          amount_annual=3300, holders="Alex Sample")
    items = ofa_store.worklist(engine, D(2026, 10, 10))
    st = items[0]["status"]
    assert st.state == "covered" and st.anniversary == D(2027, 3, 1) and not st.problems
    assert ofa_store.worklist(engine, D(2027, 2, 1))[0]["status"].state == "open"
    with pytest.raises(ofa_store.OfaError, match="already"):
        ofa_store.record_consent(engine, aid, D(2026, 3, 5), "paper", "t")
    counts = ofa_store.write_worklist(ofa_store.worklist(engine, D(2027, 7, 20)),
                                      tmp_path / "w.xlsx", D(2027, 7, 20))
    assert counts == {"urgent": 1}
    import openpyxl
    ws = openpyxl.load_workbook(tmp_path / "w.xlsx").active
    assert ws["A2"].value == HOUSEHOLD["name"] and "29 Jul 2027" in ws["C2"].value


def test_lapse_and_withdrawal_trail(engine):
    ofa_store.seed(engine, "t", today=D(2026, 10, 10))
    aid = f"OFA-{HOUSEHOLD['id']}-1"
    ofa_store.confirm_history(engine, aid, "pat", None)
    # confirmed: no consent before Brightly, so the 1 March 2025 window (closed 29 July 2025)
    with pytest.raises(ofa_store.OfaError, match="Not lapsed"):
        ofa_store.mark_lapsed(engine, aid, "pat", D(2025, 6, 1))
    s = ofa_store.worklist(engine, D(2025, 8, 3))[0]["status"]
    assert s.state == "lapsed" and s.due_on == D(2025, 8, 12)
    ofa_store.mark_lapsed(engine, aid, "pat", D(2025, 8, 3), providers_notified_on=D(2025, 8, 5))
    assert ofa_store.worklist(engine, D(2025, 8, 6))[0]["status"].state == "ended"
    with pytest.raises(ofa_store.OfaError, match="lapsed"):
        ofa_store.record_consent(engine, aid, D(2027, 2, 1), "esign", "t")
    with engine.connect() as c:
        kinds = [r[0] for r in c.execute(sa.select(db.ofa_event.c.kind)
                                         .where(db.ofa_event.c.arrangement_id == aid))]
    assert kinds[-2:] == ["lapsed", "providers_notified"]
    with pytest.raises(sa.exc.DatabaseError):            # the evidence trail can't be edited
        with engine.begin() as c:
            c.execute(db.ofa_event.delete())


def test_withdrawal(engine):
    ofa_store.seed(engine, "t", today=D(2026, 10, 10))
    aid = f"OFA-{HOUSEHOLD['id']}-1"
    due = ofa_store.record_withdrawal(engine, aid, D(2026, 10, 12), "pat", "Email from client")
    assert due["acknowledge_by"] == D(2026, 10, 26)
    assert ofa_store.worklist(engine, D(2026, 10, 13))[0]["status"].state == "ended"


def test_consent_form(engine, tmp_path):
    from docx import Document
    ofa_store.seed(engine, "t", today=D(2026, 10, 10))
    aid = f"OFA-{HOUSEHOLD['id']}-1"
    ofa_store.add_account(engine, aid, "Sample Wrap", "pat", account_number="W2",
                          amount_annual=3300, holders="Alex Sample, Sam Sample")
    form = form_for(engine, aid, D(2027, 3, 1))
    build(form, Practice(name="Example Advice", abn="12 345 678 901", phone="03 9000 0000"),
          tmp_path / "f.docx")
    d = Document(tmp_path / "f.docx")
    text = "\n".join(p.text for p in d.paragraphs)
    cells = " ".join(c.text for t in d.tables for row in t.rows for c in row.cells)
    for required in ("Why we are asking", "Example Advice", "Services you will receive",
                     "$3,300.00 a year", "end this arrangement", "29 July 2027",
                     "lasts until 29 July 2028"):
        assert required in text, required
    assert "Alex Sample" in cells and "Sam Sample" in cells and "W2" in cells
