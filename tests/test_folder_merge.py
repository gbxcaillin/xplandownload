import csv

import openpyxl

from xplan_extract.folder_merge import (DECISION, Person, apply_merges, compare, confidence,
                                        find_duplicates, group_folders, load_folders)

HEAD = ["kind", "client_id", "client_name", "other_client_ids", "date", "type", "subtype",
        "subject", "xplan_docid", "xplan_partid", "original_filename", "mimetype", "bytes",
        "saved_as", "status"]


def _setup(tmp_path):
    files = [  # client id, folder, file name, docid, partid
        ("41479", "Active\\Alexander Mcdonough (41479)", "a.pdf", "1", "1"),
        ("41479", "Active\\Alexander Mcdonough (41479)", "b.pdf", "2", "1"),
        ("40975", "Inactive\\Alexander Mc Donough (40975)", "a.pdf", "1", "1"),   # same doc
        ("40975", "Inactive\\Alexander Mc Donough (40975)", "b.pdf", "9", "1"),   # different
        ("40975", "Inactive\\Alexander Mc Donough (40975)", "c.pdf", "3", "1"),
        ("4034", "Inactive\\Alister Pillar (4034)", "x.pdf", "5", "1"),
        ("4036", "Inactive\\Alister Pillar (4036)", "y.pdf", "6", "1"),
        ("700", "Inactive\\Jane Citizen (700)", "z.pdf", "7", "1"),
    ]
    rows = []
    for cid, folder, name, docid, partid in files:
        path = tmp_path / "Clients" / folder.replace("\\", "/")
        path.mkdir(parents=True, exist_ok=True)
        (path / name).write_text(f"{folder}/{name}")
        status = folder.split("\\")[0]
        rows.append(["attachment", cid, "", "", "", "", "", "", docid, partid, name, "", "1",
                     f"Clients\\{folder}\\{name}", status])
    with open(tmp_path / "documents_index.csv", "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(HEAD)
        w.writerows(rows)


def test_compare_and_confidence():
    a = Person(1, "individual", "alexander", "", "mcdonough", "1960-01-02", {"a@x.com"})
    b = Person(2, "individual", "alex", "", "mcdonough", "1960-01-02", {"a@x.com"})
    c = Person(3, "individual", "alexander", "", "mcdonough", "1990-05-05")
    d = Person(4, "trust")
    people = {1: a, 2: b, 3: c, 4: d}
    ev, con = compare([1], [2], people, 0)
    assert "same date of birth" in ev and "same email" in ev and not con
    assert confidence(ev, con) == ("High", "merge")
    ev, con = compare([1], [3], people, 2)
    assert "different dates of birth" in con
    assert confidence(ev, con) == ("Different people?", "keep")
    assert confidence([], []) == ("Name only", "merge")
    _, con = compare([1], [4], people, 0)
    assert any("trust" in c for c in con)
    twin = Person(5, "individual", "mary", "", "mcdonough", "1960-01-02")
    _, con = compare([1], [5], {1: a, 5: twin}, 0)
    assert "different first names" in con


def test_grouping(tmp_path):
    _setup(tmp_path)
    folders = load_folders(tmp_path)
    groups = group_folders(folders)
    names = sorted(sorted(f.name for f in g) for g in groups)
    assert names == [["Alexander Mc Donough (40975)", "Alexander Mcdonough (41479)"],
                     ["Alister Pillar (4034)", "Alister Pillar (4036)"]]
    # surname + date of birth also groups different spellings of a first name
    people = {700: Person(700, "individual", "jane", "", "citizen", "1970-01-01"),
              4034: Person(4034, "individual", "alister", "", "citizen", "1970-01-01")}
    assert any({f.name for f in g} >= {"Jane Citizen (700)", "Alister Pillar (4034)"}
               for g in group_folders(folders, people))


def test_review_then_apply(tmp_path):
    _setup(tmp_path)
    proposals, review = find_duplicates(tmp_path, None, lambda m: None)
    assert len(proposals) == 2
    mc = next(p for p in proposals if "Mc" in p.main.name)
    assert mc.main.name == "Alexander Mcdonough (41479)"     # Active wins
    assert mc.shared == 1 and mc.decision == "merge"

    # mark the Pillar pair as different people
    wb = openpyxl.load_workbook(review)
    ws = wb.active
    head = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[head.index("Main folder")].value and "Pillar" in str(
                row[head.index("Main folder")].value):
            row[head.index(DECISION)].value = "keep"
    wb.save(review)

    r = apply_merges(tmp_path, review, lambda m: None)
    assert r.merged == 1 and r.kept == 1 and not r.problems
    assert r.duplicates_removed == 1 and r.renamed == 1 and r.files_moved == 2
    main = tmp_path / "Clients" / "Active" / "Alexander Mcdonough (41479)"
    assert sorted(p.name for p in main.iterdir()) == ["a.pdf", "b (from 40975).pdf", "b.pdf",
                                                      "c.pdf"]
    assert (main / "b.pdf").read_text().endswith("41479)/b.pdf")    # never overwritten
    assert not (tmp_path / "Clients" / "Inactive" / "Alexander Mc Donough (40975)").exists()
    assert (tmp_path / "Clients" / "Inactive" / "Alister Pillar (4036)").is_dir()
    with open(tmp_path / "documents_index.csv", encoding="utf-8-sig") as fh:
        rows = [r for r in csv.DictReader(fh) if r["client_id"] == "40975"]
    assert {r["saved_as"] for r in rows} == {
        "Clients\\Active\\Alexander Mcdonough (41479)\\a.pdf",
        "Clients\\Active\\Alexander Mcdonough (41479)\\b (from 40975).pdf",
        "Clients\\Active\\Alexander Mcdonough (41479)\\c.pdf"}
    assert {r["status"] for r in rows} == {"Active"}
    assert r.log_file.exists()
    # applying again finds nothing left to merge
    again = apply_merges(tmp_path, review, lambda m: None)
    assert again.merged == 0 and again.already == 1 and not again.problems


def test_apply_only_high(tmp_path):
    _setup(tmp_path)
    _, review = find_duplicates(tmp_path, None, lambda m: None)
    wb = openpyxl.load_workbook(review)
    ws = wb.active
    head = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[head.index("Main folder")].value and "Mc" in str(
                row[head.index("Main folder")].value):
            row[head.index("Confidence")].value = "High"
    wb.save(review)
    r = apply_merges(tmp_path, review, lambda m: None, only=["high"])
    assert r.merged == 1 and r.skipped == 1
    assert (tmp_path / "Clients" / "Inactive" / "Alister Pillar (4036)").is_dir()
    # the remade list only has what's left
    proposals, _ = find_duplicates(tmp_path, None, lambda m: None)
    assert [p.other.name for p in proposals] == ["Alister Pillar (4036)"]


def test_empty_leftovers_are_skipped_and_removed(tmp_path, monkeypatch):
    from xplan_extract import folder_merge
    _setup(tmp_path)
    # a merge whose folder removal failed (OneDrive holding it) leaves an empty folder
    src = tmp_path / "Clients" / "Inactive" / "Alexander Mc Donough (40975)"
    for f in src.iterdir():
        f.unlink()
    proposals, review = find_duplicates(tmp_path, None, lambda m: None)
    assert all("40975" not in p.other.name and "40975" not in p.main.name for p in proposals)
    r = apply_merges(tmp_path, review, lambda m: None, only=["nothing"])
    assert r.empty_removed == 1 and not src.exists()


def test_remove_empty_retries(tmp_path, monkeypatch):
    from xplan_extract import folder_merge
    d = tmp_path / "x"
    (d / "sub").mkdir(parents=True)
    calls = {"n": 0}
    real = folder_merge.os.rmdir

    def flaky(path):
        calls["n"] += 1
        if calls["n"] <= 2:
            raise PermissionError(13, "in use")
        real(path)

    monkeypatch.setattr(folder_merge.os, "rmdir", flaky)
    monkeypatch.setattr(folder_merge.time, "sleep", lambda s: None)
    assert folder_merge.remove_empty(d) == "" and not d.exists()
