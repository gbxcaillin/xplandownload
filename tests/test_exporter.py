import datetime as dt
import json
from decimal import Decimal

import openpyxl
import pytest
import sqlalchemy as sa

from xplan_extract import exporter
from xplan_extract.exporter import ExportOptions, export_database, to_json_value


@pytest.fixture
def engine(tmp_path):
    eng = sa.create_engine(f"sqlite:///{tmp_path / 'src.db'}")
    with eng.begin() as c:
        c.exec_driver_sql("CREATE TABLE client (id INTEGER, name TEXT, joined TEXT, photo BLOB)")
        c.exec_driver_sql(
            "INSERT INTO client VALUES (1, '=1+1', '2024-01-02', x'0102'), "
            "(2, ?, NULL, NULL)", (("y" * 40_000),)
        )
        c.exec_driver_sql("CREATE TABLE empty (id INTEGER)")
        c.exec_driver_sql("CREATE TABLE big (id INTEGER)")
        c.exec_driver_sql(
            "WITH RECURSIVE n(i) AS (SELECT 1 UNION ALL SELECT i + 1 FROM n WHERE i < 25) "
            "INSERT INTO big SELECT i FROM n"
        )
    return eng


def test_json_values():
    assert to_json_value(Decimal("1.50")) == 1.5
    assert to_json_value(Decimal("10.000")) == 10
    assert to_json_value(Decimal("12345678901234567890.123456789")) == "12345678901234567890.123456789"
    assert to_json_value(dt.datetime(2024, 1, 2, 3, 4, 5)) == "2024-01-02T03:04:05"
    assert to_json_value(b"\x01\x02") == "AQI="
    assert to_json_value(float("nan")) is None


def test_export_single_workbook(engine, tmp_path):
    out = tmp_path / "out"
    opts = ExportOptions(max_rows_per_sheet=10, batch_size=7, json_layout="both")
    manifest = export_database(engine, out, "src", opts, progress=lambda m: None)

    assert manifest["table_count"] == 3
    assert manifest["total_rows"] == 27
    tables = {t["name"]: t for t in manifest["tables"]}
    assert tables["big"]["excel_sheets"] == ["big", "big (2)", "big (3)"]
    assert tables["empty"]["json_file"] is None          # empty tables skipped by default
    assert tables["client"]["excel_truncated_cells"] == 1

    rows = json.loads((out / "json" / "client.json").read_text(encoding="utf-8"))
    assert rows[0] == {"id": 1, "name": "=1+1", "joined": "2024-01-02", "photo": "AQI="}
    assert len(rows[1]["name"]) == 40_000

    combined = json.loads((out / "json" / "src.json").read_text(encoding="utf-8"))
    assert set(combined["tables"]) == {"client", "big"}
    assert len(combined["tables"]["big"]) == 25

    wb = openpyxl.load_workbook(out / "excel" / "src.xlsx")
    assert wb.sheetnames == ["Index", "big", "big (2)", "big (3)", "client"]
    assert [r[0] for r in wb["big (3)"].iter_rows(min_row=2, values_only=True)] == [21, 22, 23, 24, 25]
    client = list(wb["client"].iter_rows(values_only=True))
    assert client[1][1] == "=1+1"                        # stored as text, not a formula
    assert len(client[2][1]) == exporter.EXCEL_MAX_CELL_CHARS
    assert client[1][3] == "0x0102"


def test_export_per_table_and_jsonl(engine, tmp_path):
    out = tmp_path / "out"
    opts = ExportOptions(excel_layout="per-table", jsonl=True, include_empty=True)
    manifest = export_database(engine, out, "src", opts, progress=lambda m: None)
    assert manifest["excel_workbook"] is None
    assert sorted(p.name for p in (out / "excel").iterdir()) == ["big.xlsx", "client.xlsx", "empty.xlsx"]
    lines = (out / "json" / "big.jsonl").read_text().splitlines()
    assert len(lines) == 25 and json.loads(lines[0]) == {"id": 1}
    assert (out / "json" / "empty.jsonl").read_text() == ""


def test_table_filter(engine, tmp_path):
    manifest = export_database(engine, tmp_path / "o", "src", ExportOptions(tables=["CLIENT"]),
                               progress=lambda m: None)
    assert [t["name"] for t in manifest["tables"]] == ["client"]


def test_unique_names():
    names = exporter.UniqueNames(31)
    long = "x" * 40
    assert names.take(long) == "x" * 31
    assert names.take(long.upper()) == "X" * 29 + "~2"
    assert names.take("Index") == "Index"
    assert names.take("index") == "index~2"
